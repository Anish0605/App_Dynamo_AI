# analysis.py — Dynamo AI (FINAL, SAFE, STRUCTURED, UI-FRIENDLY)

import io
import base64
import mimetypes

import pandas as pd
import matplotlib
matplotlib.use("Agg")  # REQUIRED for server environments
import matplotlib.pyplot as plt

from pypdf import PdfReader
from docx import Document
from google import genai
from google.genai import types as genai_types
import config

_client = genai.Client(api_key=config.GEMINI_KEY) if config.GEMINI_KEY else None


# --------------------------------------------------
# UNIVERSAL FILE ANALYSIS ENGINE
# --------------------------------------------------

def process_file_universally(file_bytes: bytes, filename: str):
    fn = filename.lower()

    try:
        # ==================================================
        # 1️⃣ TABULAR DATA (CSV / EXCEL)
        # ==================================================
        if fn.endswith((".csv", ".xlsx", ".xls")):

            try:
                if fn.endswith(".csv"):
                    df = pd.read_csv(
                        io.BytesIO(file_bytes),
                        encoding="utf-8",
                        errors="ignore"
                    )
                else:
                    df = pd.read_excel(io.BytesIO(file_bytes))
            except Exception:
                return {
                    "type": "text",
                    "content": "Unable to read tabular data.",
                    "insight": "File format not supported or corrupted."
                }

            # 🔒 Clean & normalize dataframe
            df = df.fillna("").astype(str)

            columns = list(df.columns)
            rows = df.head(10).values.tolist()

            # Detect numeric columns (for chart)
            numeric_df = df.apply(pd.to_numeric, errors="coerce")
            numeric_df = numeric_df.dropna(axis=1, how="all")

            # -------------------------------
            # 📊 Chart + Table
            # -------------------------------
            if not numeric_df.empty:
                plt.figure(figsize=(10, 5))
                numeric_df.head(10).plot(kind="bar", color="#EAB308")
                plt.title(f"Dynamo Analysis: {filename}")
                plt.tight_layout()

                buf = io.BytesIO()
                plt.savefig(buf, format="png")
                plt.close()
                buf.seek(0)

                img_b64 = base64.b64encode(buf.read()).decode()

                return {
                    "type": "chart",
                    "image": "data:image/png;base64," + img_b64,
                    "columns": columns,
                    "rows": rows,
                    "insight": f"Extracted numeric trends from {filename}. Showing first 10 rows."
                }

            # -------------------------------
            # 📋 Table only
            # -------------------------------
            return {
                "type": "table",
                "columns": columns,
                "rows": rows,
                "insight": f"Preview of first 10 rows from {filename}. No numeric columns detected."
            }

        # ==================================================
        # 2️⃣ DOCUMENTS (PDF / DOCX / TXT)
        # ==================================================
        elif fn.endswith((".pdf", ".docx", ".txt")):
            text = ""

            if fn.endswith(".pdf"):
                reader = PdfReader(io.BytesIO(file_bytes))
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"

            elif fn.endswith(".docx"):
                doc = Document(io.BytesIO(file_bytes))
                for para in doc.paragraphs:
                    if para.text:
                        text += para.text + "\n"

            else:  # TXT
                text = file_bytes.decode("utf-8", errors="ignore")

            return {
                "type": "text",
                "content": text[:30000],  # token safety
                "insight": f"Read {filename} successfully."
            }

        # ==================================================
        # 3️⃣ IMAGE / VISION ANALYSIS
        # ==================================================
        elif fn.endswith((".png", ".jpg", ".jpeg", ".webp")):

            if not config.GEMINI_KEY:
                return {
                    "type": "text",
                    "content": "Vision analysis is not configured."
                }

            mime_type = mimetypes.guess_type(filename)[0] or "image/png"
            img_b64 = base64.b64encode(file_bytes).decode()

            response = _client.models.generate_content(
                model="gemini-3.5-flash-lite",
                contents=[
                    genai_types.Part.from_bytes(data=file_bytes, mime_type=mime_type),
                    "Describe this image for research purposes."
                ]
            )

            return {
                "type": "vision",
                "content": response.text,
                "image": f"data:{mime_type};base64,{img_b64}",
                "insight": "Visual analysis complete."
            }

    except Exception as e:
        return {
            "type": "text",
            "content": "Analysis failed.",
            "error": str(e)
        }

    # --------------------------------------------------
    # Unsupported file
    # --------------------------------------------------
    return {
        "type": "text",
        "content": "Unsupported file format."
    }


# --------------------------------------------------
# DEEP SPREADSHEET ANALYSIS (Data Analysis feature)
# Completely separate from process_file_universally.
# Called only by /data-analysis-chat endpoint.
# --------------------------------------------------

def analyze_spreadsheet_deep(file_bytes: bytes, filename: str) -> dict:
    """
    Parse a CSV/Excel file, produce a 4-panel matplotlib chart
    (top performers, bottom performers, histogram, box-plot),
    compute rich statistics including quartiles/IQR/outlier detection,
    and return everything for the /data-analysis-chat endpoint.
    Does NOT touch process_file_universally or any existing function.
    """
    fn = filename.lower()

    # ── Keywords that appear in real header rows ────────────────────
    _HEADER_KW = {
        "symbol", "isin", "quantity", "profit", "loss", "pnl", "p&l",
        "date", "value", "trade", "exchange", "product", "realized",
        "amount", "price", "buy", "sell", "return", "net", "total",
        "revenue", "cost", "income", "score", "name", "category",
    }

    def _header_score(row_series: pd.Series) -> int:
        """Count how many cells in a row match known header keywords."""
        cells = row_series.fillna("").astype(str).str.lower().tolist()
        return sum(1 for c in cells if any(kw in c for kw in _HEADER_KW))

    def _find_header_row(raw_df: pd.DataFrame, max_scan: int = 30) -> int:
        """Return the row index (0-based) that looks most like a header row."""
        best_row, best_score = 0, 0
        for i in range(min(max_scan, len(raw_df))):
            score = _header_score(raw_df.iloc[i])
            if score > best_score:
                best_score, best_row = score, i
        return best_row if best_score >= 2 else 0

    def _clean_df(d: pd.DataFrame) -> pd.DataFrame:
        """Drop all-empty rows/cols; drop rows whose first real column equals a
        known section-header keyword (handles multi-section Zerodha reports)."""
        d = d.dropna(how="all").reset_index(drop=True)
        # Drop columns that are entirely empty or all-unnamed with no real data
        d = d.dropna(axis=1, how="all")
        unnamed_empty = [
            c for c in d.columns
            if str(c).startswith("Unnamed:")
            and d[c].astype(str).str.strip().replace("nan", "").eq("").all()
        ]
        if unnamed_empty:
            d = d.drop(columns=unnamed_empty)
        if len(d) == 0:
            return d
        first_col = d.columns[0]
        # Remove rows where first cell is a repeated column header or section label
        header_vals = set(str(c).strip().lower() for c in d.columns)
        mask = d[first_col].astype(str).str.strip().str.lower().apply(
            lambda v: v not in header_vals and v not in ("", "nan", "total", "subtotal", "equity", "f&o", "futures", "options")
        )
        return d[mask].reset_index(drop=True)

    def _unnamed_ratio(d: pd.DataFrame) -> float:
        return sum(1 for c in d.columns if str(c).startswith("Unnamed:")) / max(len(d.columns), 1)

    def _read_smart(raw_df: pd.DataFrame, reload_fn) -> pd.DataFrame:
        """Auto-detect header row only when needed (unnamed columns present),
        otherwise trust pandas default read. Avoids false positives on clean CSVs."""
        default_df = _clean_df(reload_fn(0))
        # If pandas parsed it fine (< 30% unnamed), use it as-is
        if _unnamed_ratio(default_df) < 0.30:
            return default_df
        # Otherwise scan for the real header row (brokerage/metadata-prefixed files)
        header_row = _find_header_row(raw_df)
        if header_row == 0:
            return default_df
        df_out = reload_fn(header_row)
        return _clean_df(df_out)

    try:
        if fn.endswith(".csv"):
            enc = "utf-8"
            try:
                raw_df = pd.read_csv(io.BytesIO(file_bytes), encoding=enc,
                                     header=None, on_bad_lines="skip")
            except UnicodeDecodeError:
                enc = "latin-1"
                raw_df = pd.read_csv(io.BytesIO(file_bytes), encoding=enc,
                                     header=None, on_bad_lines="skip")

            def _reload_csv(skip):
                return pd.read_csv(io.BytesIO(file_bytes), encoding=enc,
                                   skiprows=skip, on_bad_lines="skip")

            df = _read_smart(raw_df, _reload_csv)

        elif fn.endswith((".xlsx", ".xls")):
            raw_df = pd.read_excel(io.BytesIO(file_bytes), header=None)

            def _reload_excel(skip):
                return pd.read_excel(io.BytesIO(file_bytes), skiprows=skip)

            df = _read_smart(raw_df, _reload_excel)
        else:
            return {"error": "Please upload a .csv or .xlsx / .xls file."}
    except Exception as e:
        return {"error": f"Could not read file: {str(e)[:120]}"}

    if df.empty:
        return {"error": "File appears to be empty."}

    row_count = len(df)

    # ── Column classification ───────────────────────────────────────
    numeric_df = df.apply(pd.to_numeric, errors="coerce")
    all_numeric_cols = [c for c in numeric_df.columns if not numeric_df[c].isna().all()]

    label_col = None
    for col in df.columns:
        if numeric_df[col].isna().all():
            label_col = col
            break

    value_keywords = [
        "profit", "pnl", "p&l", "gain", "loss", "return",
        "realized", "net", "amount", "value", "total", "revenue",
        "sales", "price", "cost", "income", "expense", "score",
    ]
    value_col = None
    for kw in value_keywords:
        for col in all_numeric_cols:
            if kw in str(col).lower():
                value_col = col
                break
        if value_col:
            break
    if value_col is None and all_numeric_cols:
        value_col = all_numeric_cols[0]

    # ── Build sorted frame + full statistics ───────────────────────
    df_clean = None
    vals_series = None
    stats = {}
    outliers = []

    if value_col:
        tmp = df[[label_col, value_col]].copy() if label_col else df[[value_col]].copy()
        tmp["_val"] = pd.to_numeric(tmp[value_col], errors="coerce")
        tmp = tmp.dropna(subset=["_val"])
        if label_col:
            tmp = tmp.rename(columns={label_col: "label"})
            tmp["value"] = tmp["_val"]
        else:
            tmp["label"] = tmp.index.astype(str)
            tmp["value"] = tmp["_val"]
        df_clean = tmp[["label", "value"]].sort_values("value", ascending=False).reset_index(drop=True)
        vals_series = df_clean["value"]

        q1      = float(vals_series.quantile(0.25))
        q3      = float(vals_series.quantile(0.75))
        iqr     = q3 - q1
        pos     = int((vals_series > 0).sum())
        neg     = int((vals_series < 0).sum())
        wr      = pos / len(vals_series) * 100 if len(vals_series) > 0 else 0
        std_dev = float(vals_series.std()) if len(vals_series) > 1 else 0.0

        stats = {
            "count":    len(vals_series),
            "sum":      float(vals_series.sum()),
            "mean":     float(vals_series.mean()),
            "median":   float(vals_series.median()),
            "std":      std_dev,
            "min":      float(vals_series.min()),
            "max":      float(vals_series.max()),
            "q1":       q1,
            "q3":       q3,
            "iqr":      iqr,
            "pos":      pos,
            "neg":      neg,
            "win_rate": wr,
        }

        lower = q1 - 1.5 * iqr
        upper = q3 + 1.5 * iqr
        mask  = (vals_series < lower) | (vals_series > upper)
        if mask.any():
            outliers = df_clean[mask.values].head(10).to_dict("records")

    # ── 4-panel dark chart ──────────────────────────────────────────
    chart_b64 = None
    if df_clean is not None and vals_series is not None and len(df_clean) >= 2:
        n        = len(df_clean)
        show     = min(10, n)
        top_n    = df_clean.head(show)
        bottom_n = df_clean.tail(show).sort_values("value")

        fig, axes = plt.subplots(2, 2, figsize=(14, 10))
        fig.patch.set_facecolor("#111827")

        def _style(ax):
            ax.set_facecolor("#1f2937")
            ax.tick_params(colors="#9ca3af", labelsize=8)
            for sp in ax.spines.values():
                sp.set_color("#374151")

        # Panel 1 — Top performers
        ax1 = axes[0, 0]
        _style(ax1)
        c1 = ["#22c55e" if v >= 0 else "#f59e0b" for v in top_n["value"].values[::-1]]
        ax1.barh(top_n["label"].astype(str)[::-1], top_n["value"].values[::-1],
                 color=c1, edgecolor="none", height=0.65)
        ax1.axvline(0, color="#4b5563", linewidth=0.8)
        ax1.set_title(f"Top Performers — {value_col}", color="#f9fafb",
                      fontsize=10, fontweight="bold", pad=8)

        # Panel 2 — Bottom performers
        ax2 = axes[0, 1]
        _style(ax2)
        c2 = ["#ef4444" if v < 0 else "#f59e0b" for v in bottom_n["value"].values]
        ax2.barh(bottom_n["label"].astype(str), bottom_n["value"].values,
                 color=c2, edgecolor="none", height=0.65)
        ax2.axvline(0, color="#4b5563", linewidth=0.8)
        ax2.set_title(f"Bottom Performers — {value_col}", color="#f9fafb",
                      fontsize=10, fontweight="bold", pad=8)

        # Panel 3 — Distribution histogram with mean/median lines
        ax3 = axes[1, 0]
        _style(ax3)
        bins = min(25, max(8, n // 3))
        ax3.hist(vals_series.dropna(), bins=bins, color="#3b82f6", edgecolor="#1e40af", alpha=0.85)
        mean_v   = stats.get("mean", 0)
        median_v = stats.get("median", 0)
        ax3.axvline(mean_v,   color="#eab308", linewidth=1.8, linestyle="--",
                    label=f"Mean {mean_v:.1f}")
        ax3.axvline(median_v, color="#22c55e", linewidth=1.8, linestyle=":",
                    label=f"Median {median_v:.1f}")
        ax3.axvline(0, color="#4b5563", linewidth=0.8)
        ax3.legend(fontsize=7, facecolor="#374151", labelcolor="#d1d5db", framealpha=0.9)
        ax3.set_title(f"Distribution of {value_col}", color="#f9fafb",
                      fontsize=10, fontweight="bold", pad=8)

        # Panel 4 — Box plot (quartiles + outliers)
        ax4 = axes[1, 1]
        _style(ax4)
        clean_vals = vals_series.dropna().tolist()
        if len(clean_vals) >= 4:
            ax4.boxplot(
                clean_vals,
                patch_artist=True,
                whiskerprops=dict(color="#9ca3af", linewidth=1.5),
                capprops=dict(color="#9ca3af", linewidth=1.5),
                medianprops=dict(color="#eab308", linewidth=2.5),
                flierprops=dict(marker="o", color="#ef4444", markersize=5, alpha=0.8),
                boxprops=dict(facecolor="#3b82f6", alpha=0.45, linewidth=1.5, edgecolor="#6b7280"),
            )
            ax4.set_xticks([])
            q1v = stats.get("q1", 0)
            q3v = stats.get("q3", 0)
            title_str = (
                f"Box Plot   "
                f"Q1={q1v:.1f}  |  Med={median_v:.1f}  |  Q3={q3v:.1f}"
            )
            ax4.set_title(title_str, color="#f9fafb", fontsize=9, fontweight="bold", pad=8)
        else:
            ax4.scatter(range(len(clean_vals)), clean_vals, color="#3b82f6", s=60, alpha=0.8)
            ax4.set_title(f"All Values — {value_col}", color="#f9fafb",
                          fontsize=10, fontweight="bold", pad=8)

        plt.tight_layout(pad=2.0)
        buf = io.BytesIO()
        plt.savefig(buf, format="png", dpi=120, facecolor=fig.get_facecolor())
        plt.close()
        buf.seek(0)
        chart_b64 = "data:image/png;base64," + base64.b64encode(buf.read()).decode()

    # ── Gemini-readable summary (all real numbers, no placeholders) ─
    preview_cols = [str(c) for c in df.columns.tolist()]
    preview_rows = df.head(5).fillna("").astype(str).values.tolist()

    tbl_md = "| " + " | ".join(preview_cols) + " |\n"
    tbl_md += "| " + " | ".join(["---"] * len(preview_cols)) + " |\n"
    for row in preview_rows:
        tbl_md += "| " + " | ".join(str(v)[:40] for v in row) + " |\n"

    stats_block = ""
    if df_clean is not None and vals_series is not None:
        top10_str = df_clean.head(10).to_string(index=False)
        bot10_str = df_clean.tail(10).sort_values("value").to_string(index=False)

        outlier_str = ""
        if outliers:
            outlier_str = f"\nOutliers (IQR method) — {len(outliers)} detected:\n"
            for o in outliers[:8]:
                lbl = o.get("label", "?")
                val = o.get("value", 0)
                outlier_str += f"  {lbl}: {val:.2f}\n"

        stats_block = f"""
Comprehensive statistics for column '{value_col}':
  Count:        {stats['count']}
  Sum:          {stats['sum']:.2f}
  Mean:         {stats['mean']:.2f}
  Median:       {stats['median']:.2f}
  Std Dev:      {stats['std']:.2f}
  Min:          {stats['min']:.2f}
  Max:          {stats['max']:.2f}
  Q1 (25th %):  {stats['q1']:.2f}
  Q3 (75th %):  {stats['q3']:.2f}
  IQR:          {stats['iqr']:.2f}
  Positive:     {stats['pos']}  ({stats['win_rate']:.1f}% win rate)
  Negative:     {stats['neg']}
  Outliers:     {len(outliers)}
{outlier_str}
Top 10 by {value_col}:
{top10_str}

Bottom 10 by {value_col}:
{bot10_str}
"""
        if len(all_numeric_cols) > 1:
            stats_block += f"\nAll numeric columns in file: {', '.join(all_numeric_cols)}\n"

    summary_text = (
        f"File: {filename}  |  {row_count} rows  |  {len(df.columns)} columns\n"
        f"Columns: {', '.join(preview_cols)}\n\n"
        f"First 5 rows:\n{tbl_md}\n"
        f"{stats_block}"
    )

    table_preview = {
        "columns": preview_cols,
        "rows": df.head(10).fillna("").astype(str).values.tolist(),
    }

    csv_str = df_clean.to_csv(index=False) if df_clean is not None else df.head(200).fillna("").to_csv(index=False)
    csv_b64 = base64.b64encode(csv_str.encode()).decode()

    return {
        "chart_b64":         chart_b64,
        "summary_text":      summary_text,
        "table_preview":     table_preview,
        "download_csv_b64":  csv_b64,
        "row_count":         row_count,
        "value_col":         value_col or "",
        "filename":          filename,
        "stats":             stats,
        "outliers":          outliers,
        "all_numeric_cols":  all_numeric_cols,
    }
